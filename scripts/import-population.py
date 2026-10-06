"""Import the supplied 2026 workbook without modifying existing election data.

Usage: python scripts/import-population.py /path/to/Fylker_og_kommuner_2026.xlsx
Requires openpyxl for read-only workbook extraction.
"""
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "dist" / "data"


def import_population(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    municipalities = {}
    counties = {}
    totals = defaultdict(int)
    counts = defaultdict(int)
    try:
        for code, name, county, county_name, population, *_ in workbook["Kommuner"].iter_rows(min_row=2, values_only=True):
            if code is None or not str(code).isdigit():
                continue
            code, county = str(code).zfill(4), str(county).zfill(2)
            assert code not in municipalities, f"Duplicate municipality: {code}"
            assert isinstance(population, int) and population > 0, f"Invalid population: {code}"
            municipalities[code] = population
            totals[county] += population
            counts[county] += 1
        for code, name, count, population, share, control, difference in workbook["Fylker"].iter_rows(min_row=2, values_only=True):
            if code is None or not str(code).isdigit():
                continue
            code = str(code).zfill(2)
            assert code not in counties, f"Duplicate county: {code}"
            assert population == totals[code] == control and difference == 0, f"County total mismatch: {code}"
            assert count == counts[code], f"Municipality count mismatch: {code}"
            counties[code] = population
    finally:
        workbook.close()
    existing_municipalities = json.loads((DATA / "municipalities.json").read_text(encoding="utf-8"))
    existing_counties = json.loads((DATA / "counties.json").read_text(encoding="utf-8"))
    assert set(municipalities) == {m["id"] for m in existing_municipalities}, "Municipality coverage differs"
    assert set(counties) == {c["code"] for c in existing_counties}, "County coverage differs"
    assert sum(municipalities.values()) == sum(counties.values()), "National totals differ"
    result = {
        "asOf": "2026-01-01",
        "source": "https://www.ssb.no/statbank/table/07459/",
        "sourceLabel": "SSB tabell 07459",
        "sourceWorkbook": Path(path).name,
        "municipalities": dict(sorted(municipalities.items())),
        "counties": dict(sorted(counties.items())),
    }
    (DATA / "population.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(municipalities)} municipalities and {len(counties)} counties; total {sum(counties.values()):,}.")
    print(f"Bergen: {municipalities['4601']:,}; Vestland: {counties['46']:,}.")


if __name__ == "__main__":
    import_population(sys.argv[1])
